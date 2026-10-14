"""
generate_test_artifacts.py - Generates sample .docx and scanned image PDF files for verification.
"""
from pathlib import Path
import docx
from docx import Document
from PIL import Image, ImageDraw, ImageFont
import fitz

DATA_DIR = Path(__file__).resolve().parent.parent / "data" / "sample_artifacts"
DATA_DIR.mkdir(parents=True, exist_ok=True)

def generate_sample_docx():
    """Creates a sample Cadence Financial Group TPRM document in .docx format."""
    doc_path = DATA_DIR / "Cadence_TPRM_Training.docx"
    if doc_path.exists():
        return doc_path

    doc = Document()
    doc.add_heading("Cadence Financial Group - Third Party Risk Management (TPRM) Training", level=1)
    
    p1 = doc.add_paragraph("This document outlines confidential third-party oversight protocols.")
    p2 = doc.add_paragraph("For TPRM inquiries, contact lead officer Marguerite Vossberg at m.vossberg@cadencefg.example or phone +91 98765 43210.")
    p3 = doc.add_paragraph("Employee ID: EMP-90421 registered under risk division in New York.")
    
    table = doc.add_table(rows=1, cols=3)
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = "Vendor Name"
    hdr_cells[1].text = "Contact Officer"
    hdr_cells[2].text = "Direct Email"
    
    row_cells = table.add_row().cells
    row_cells[0].text = "Ironclad Solutions Inc"
    row_cells[1].text = "David Franklin"
    row_cells[2].text = "d.franklin@ironclad.example"
    
    doc.save(str(doc_path))
    print(f"Created sample DOCX: {doc_path}")
    return doc_path

def generate_scanned_pdf():
    """
    Creates a rasterized single-page PDF containing text rendered purely as an image.
    Has 0 native text characters (< 15 characters threshold), forcing OCR fallback.
    """
    pdf_path = DATA_DIR / "Cadence_Scanned_Memo.pdf"
    if pdf_path.exists():
        return pdf_path

    # Generate an image with text
    img = Image.new("RGB", (800, 600), color=(255, 255, 255))
    draw = ImageDraw.Draw(img)
    text = (
        "CADENCE FINANCIAL GROUP - SCANNED INTERNAL AUDIT MEMO\n"
        "Confidential Personnel Verification\n"
        "Officer: Johnathan Smith\n"
        "Email: j.smith@cadencefg.com\n"
        "Account ID: AC-882910\n"
        "Location: London"
    )
    draw.text((50, 50), text, fill=(0, 0, 0))

    img_temp = DATA_DIR / "temp_scanned.png"
    img.save(str(img_temp))

    # Create PDF with the image
    doc = fitz.open()
    page = doc.new_page(width=800, height=600)
    page.insert_image(page.rect, filename=str(img_temp))
    doc.save(str(pdf_path))
    doc.close()

    if img_temp.exists():
        img_temp.unlink()

    print(f"Created scanned image PDF (<15 native chars): {pdf_path}")
    return pdf_path

if __name__ == "__main__":
    generate_sample_docx()
    generate_scanned_pdf()
